"""Stage 6b, teacher analytics: monitoring board (ФТ-04), individual tasks (ФТ-05), graded
examples (ФТ-17), solve rates (ФТ-20), process metrics (ФТ-33), GenAI use (ФТ-38) and the
research export to Excel (ФТ-49).
"""

import io
from datetime import UTC, datetime
from typing import Any

from fastapi import APIRouter, HTTPException, Request, status
from fastapi.responses import Response
from openpyxl import Workbook
from pydantic import BaseModel, Field
from sqlalchemy import select

from app.api.deps import AdminUser, AppSettings, CurrentUser, DbSession, client_ip
from app.api.privacy import _assessment_for_teacher, _rows
from app.core.enums import GenAiMode, TaskStatus
from app.core.roles import Role
from app.models import (
    AiMessage,
    AppSetting,
    Course,
    CourseGroup,
    GroupMember,
    GroupTeacher,
    RubricExample,
    Submission,
    Task,
    TaskAssignment,
    User,
)
from app.services import analytics, audit
from app.services.access import can_edit_course, can_view_user, studies_course
from app.services.ai_grading import find_injection
from app.services.assignments import can_work_on_assignment
from app.services.audit import AuditAction
from app.services.contests import in_unfinished_assessment
from app.services.research import agreement
from app.services.results import confirmed_grade, latest_ai_grade

router = APIRouter(prefix="/api", tags=["insights"])


def _forbidden() -> HTTPException:
    return HTTPException(status.HTTP_403_FORBIDDEN, detail="forbidden")


def _course_for_teacher(db: DbSession, user: User, course_id: int) -> Course:
    course = db.get(Course, course_id)
    if course is None or not can_edit_course(db, user, course_id):
        raise _forbidden()
    return course


def _my_students(db: DbSession, user: User, course_id: int, group_id: int | None) -> list[User]:
    """Students of the course's groups; a teacher sees only the groups they teach."""
    groups = select(CourseGroup.group_id).where(CourseGroup.course_id == course_id)
    if not user.has_role(Role.ADMIN):
        groups = groups.where(
            CourseGroup.group_id.in_(
                select(GroupTeacher.group_id).where(GroupTeacher.user_id == user.id)
            )
        )
    if group_id is not None:
        groups = groups.where(CourseGroup.group_id == group_id)
    members = select(GroupMember.user_id).where(GroupMember.group_id.in_(groups))
    return list(db.scalars(select(User).where(User.id.in_(members)).order_by(User.full_name)))


# Monitoring board (ФТ-04) --------------------------------------------------------------------


@router.get("/courses/{course_id}/monitoring")
def monitoring(
    course_id: int, user: CurrentUser, db: DbSession, group_id: int | None = None
) -> dict[str, Any]:
    _course_for_teacher(db, user, course_id)
    students = _my_students(db, user, course_id, group_id)
    board = analytics.monitoring(db, course_id, [s.id for s in students])
    names = {s.id: s.full_name for s in students}
    for row in board["students"]:
        row["full_name"] = names[row["user_id"]]
    return board


# Individual tasks (ФТ-05) --------------------------------------------------------------------


class AssignmentIn(BaseModel):
    task_id: int
    user_id: int
    note: str = Field(default="", max_length=5000)
    genai_mode: GenAiMode = GenAiMode.HINTS_ONLY
    due_at: datetime | None = None


class AssignmentOut(BaseModel):
    id: int
    task_id: int
    task_title: str
    user_id: int
    full_name: str | None
    note: str
    genai_mode: GenAiMode
    due_at: datetime | None
    created_at: datetime
    last_verdict: str | None


def _assignment_out(db: DbSession, a: TaskAssignment, with_name: bool) -> AssignmentOut:
    task = db.get(Task, a.task_id)
    student = db.get(User, a.user_id) if with_name else None
    last = db.scalar(
        select(Submission.verdict)
        .where(
            Submission.user_id == a.user_id,
            Submission.task_id == a.task_id,
            Submission.assessment_id.is_(None),
        )
        .order_by(Submission.id.desc())
        .limit(1)
    )
    return AssignmentOut(
        id=a.id,
        task_id=a.task_id,
        task_title=task.title if task else "",
        user_id=a.user_id,
        full_name=student.full_name if student else None,
        note=a.note,
        genai_mode=GenAiMode(a.genai_mode),
        due_at=a.due_at,
        created_at=a.created_at,
        last_verdict=last,
    )


@router.post("/courses/{course_id}/assignments", status_code=status.HTTP_201_CREATED)
def create_assignment(
    course_id: int, body: AssignmentIn, request: Request, user: CurrentUser, db: DbSession
) -> AssignmentOut:
    _course_for_teacher(db, user, course_id)
    task = db.get(Task, body.task_id)
    if task is None or task.course_id != course_id or task.status != TaskStatus.APPROVED.value:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, detail="task_not_approved")
    if in_unfinished_assessment(db, task.id, datetime.now(UTC)):
        # The student would see a task of an upcoming or running control work early.
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, detail="task_in_assessment")
    if not can_view_user(db, user, body.user_id) or not studies_course(db, body.user_id, course_id):
        raise _forbidden()
    if db.scalar(
        select(TaskAssignment.id).where(
            TaskAssignment.task_id == task.id, TaskAssignment.user_id == body.user_id
        )
    ):
        raise HTTPException(status.HTTP_409_CONFLICT, detail="already_assigned")
    assignment = TaskAssignment(
        task_id=task.id,
        user_id=body.user_id,
        assigned_by=user.id,
        note=body.note,
        genai_mode=body.genai_mode.value,
        due_at=body.due_at,
    )
    db.add(assignment)
    db.flush()
    audit.record(
        db,
        AuditAction.TASK_ASSIGNED,
        actor_id=user.id,
        object_type="task_assignment",
        object_id=assignment.id,
        ip=client_ip(request),
        details={"task_id": task.id, "user_id": body.user_id},
    )
    db.commit()
    db.refresh(assignment)
    return _assignment_out(db, assignment, with_name=True)


@router.get("/courses/{course_id}/assignments")
def list_assignments(course_id: int, user: CurrentUser, db: DbSession) -> list[AssignmentOut]:
    _course_for_teacher(db, user, course_id)
    rows = db.scalars(
        select(TaskAssignment)
        .join(Task, Task.id == TaskAssignment.task_id)
        .where(Task.course_id == course_id)
        .order_by(TaskAssignment.id.desc())
    )
    return [
        _assignment_out(db, a, with_name=True) for a in rows if can_view_user(db, user, a.user_id)
    ]


@router.delete("/assignments/{assignment_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_assignment(
    assignment_id: int, request: Request, user: CurrentUser, db: DbSession
) -> None:
    assignment = db.get(TaskAssignment, assignment_id)
    task = db.get(Task, assignment.task_id) if assignment else None
    if (
        assignment is None
        or task is None
        or not can_edit_course(db, user, task.course_id)
        or not can_view_user(db, user, assignment.user_id)
    ):
        raise _forbidden()
    audit.record(
        db,
        AuditAction.TASK_UNASSIGNED,
        actor_id=user.id,
        object_type="task_assignment",
        object_id=assignment.id,
        ip=client_ip(request),
        details={"task_id": assignment.task_id, "user_id": assignment.user_id},
    )
    db.delete(assignment)
    db.commit()


@router.get("/my/assignments")
def my_assignments(user: CurrentUser, db: DbSession) -> list[AssignmentOut]:
    rows = db.scalars(
        select(TaskAssignment)
        .where(TaskAssignment.user_id == user.id)
        .order_by(TaskAssignment.id.desc())
    )
    out = []
    for a in rows:
        task = db.get(Task, a.task_id)
        if task is not None and can_work_on_assignment(db, user.id, task):
            out.append(_assignment_out(db, a, with_name=False))
    return out


# Graded works as rubric examples (ФТ-17) -----------------------------------------------------


class ExampleIn(BaseModel):
    criterion_code: str = Field(min_length=1, max_length=32)
    note: str = Field(default="", max_length=2000)


class ExampleOut(BaseModel):
    id: int
    submission_id: int
    criterion_code: str
    level: str
    note: str
    created_at: datetime


def _example_out(e: RubricExample) -> ExampleOut:
    return ExampleOut(
        id=e.id,
        submission_id=e.submission_id,
        criterion_code=e.criterion_code,
        level=e.level,
        note=e.note,
        created_at=e.created_at,
    )


@router.post("/submissions/{submission_id}/examples", status_code=status.HTTP_201_CREATED)
def add_example(
    submission_id: int, body: ExampleIn, request: Request, user: CurrentUser, db: DbSession
) -> ExampleOut:
    """The level comes from the teacher's confirmed grade, never from the AI. The code goes
    to the LLM in later gradings, so code that looks like a prompt injection is refused."""
    submission = db.get(Submission, submission_id)
    task = db.get(Task, submission.task_id) if submission else None
    if (
        submission is None
        or task is None
        or not can_edit_course(db, user, task.course_id)
        or not can_view_user(db, user, submission.user_id)
    ):
        raise _forbidden()
    ai = latest_ai_grade(db, submission.id)
    if find_injection({"code": submission.code}) or (ai and "possible_injection" in ai.flags):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, detail="possible_injection")
    grade = confirmed_grade(db, submission.id)
    level = next(
        (c["level"] for c in (grade.criteria if grade else []) if c["code"] == body.criterion_code),
        None,
    )
    if level is None:
        raise HTTPException(status.HTTP_409_CONFLICT, detail="not_confirmed")
    if db.scalar(
        select(RubricExample.id).where(
            RubricExample.submission_id == submission.id,
            RubricExample.criterion_code == body.criterion_code,
        )
    ):
        raise HTTPException(status.HTTP_409_CONFLICT, detail="example_exists")
    example = RubricExample(
        course_id=task.course_id,
        submission_id=submission.id,
        criterion_code=body.criterion_code,
        level=level,
        note=body.note,
        added_by=user.id,
    )
    db.add(example)
    db.flush()
    audit.record(
        db,
        AuditAction.RUBRIC_EXAMPLE_ADDED,
        actor_id=user.id,
        object_type="rubric_example",
        object_id=example.id,
        ip=client_ip(request),
        details={"submission_id": submission.id, "criterion": body.criterion_code},
    )
    db.commit()
    db.refresh(example)
    return _example_out(example)


@router.get("/courses/{course_id}/rubric/examples")
def list_examples(course_id: int, user: CurrentUser, db: DbSession) -> list[ExampleOut]:
    _course_for_teacher(db, user, course_id)
    rows = db.scalars(
        select(RubricExample)
        .where(RubricExample.course_id == course_id)
        .order_by(RubricExample.criterion_code, RubricExample.level, RubricExample.id)
    )
    return [_example_out(e) for e in rows]


@router.delete("/rubric-examples/{example_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_example(example_id: int, request: Request, user: CurrentUser, db: DbSession) -> None:
    example = db.get(RubricExample, example_id)
    if example is None or not can_edit_course(db, user, example.course_id):
        raise _forbidden()
    audit.record(
        db,
        AuditAction.RUBRIC_EXAMPLE_REMOVED,
        actor_id=user.id,
        object_type="rubric_example",
        object_id=example.id,
        ip=client_ip(request),
    )
    db.delete(example)
    db.commit()


# Solve rates (ФТ-20) -------------------------------------------------------------------------


@router.get("/courses/{course_id}/solve-rates")
def solve_rates(course_id: int, user: CurrentUser, db: DbSession) -> dict[str, Any]:
    _course_for_teacher(db, user, course_id)
    return {
        "settings": analytics.solve_rate_settings(db).model_dump(),
        "reference": analytics.reference_rates(),
        "tasks": analytics.solve_rates(db, course_id),
    }


@router.put("/settings/solve-rate")
def save_solve_rate_settings(
    body: analytics.SolveRateSettings, request: Request, admin: AdminUser, db: DbSession
) -> analytics.SolveRateSettings:
    audit.record(
        db,
        AuditAction.SETTING_CHANGED,
        actor_id=admin.id,
        object_type="setting",
        object_id=analytics.SOLVE_RATE_KEY,
        ip=client_ip(request),
        details=body.model_dump(),
    )
    stored = db.get(AppSetting, analytics.SOLVE_RATE_KEY)
    if stored is None:
        db.add(
            AppSetting(key=analytics.SOLVE_RATE_KEY, value=body.model_dump(), updated_by=admin.id)
        )
    else:
        stored.value = body.model_dump()
        stored.updated_by = admin.id
    db.commit()
    return body


# Process metrics (ФТ-33) and GenAI use (ФТ-38) -----------------------------------------------


def _student_rows(db: DbSession, user: User, assessment_id: int) -> list[tuple[User, Submission]]:
    """The latest submission of each student and task in the assessment."""
    assessment = _assessment_for_teacher(db, user, assessment_id)
    latest: dict[tuple[int, int], Submission] = {}
    for s in db.scalars(
        select(Submission).where(Submission.assessment_id == assessment.id).order_by(Submission.id)
    ):
        latest[(s.user_id, s.task_id)] = s
    out = []
    for (user_id, _), submission in sorted(latest.items()):
        if not can_view_user(db, user, user_id):
            continue
        student = db.get(User, user_id)
        if student is not None:
            out.append((student, submission))
    return out


@router.get("/assessments/{assessment_id}/process")
def assessment_process(
    assessment_id: int, user: CurrentUser, db: DbSession
) -> list[dict[str, Any]]:
    return [
        {
            "user_id": student.id,
            "full_name": student.full_name,
            "task_id": s.task_id,
            "submission_id": s.id,
            "verdict": s.verdict,
            **analytics.process_report(
                analytics.work_events(db, s.user_id, s.task_id, s.assessment_id)
            ),
        }
        for student, s in _student_rows(db, user, assessment_id)
    ]


@router.get("/assessments/{assessment_id}/genai")
def assessment_genai(assessment_id: int, user: CurrentUser, db: DbSession) -> dict[str, Any]:
    rows = []
    totals: dict[str, int] = {}
    for student, s in _student_rows(db, user, assessment_id):
        messages = db.scalars(
            select(AiMessage)
            .where(
                AiMessage.user_id == s.user_id,
                AiMessage.task_id == s.task_id,
                AiMessage.assessment_id == s.assessment_id,
            )
            .order_by(AiMessage.id)
        ).all()
        report = analytics.genai_report(
            analytics.work_events(db, s.user_id, s.task_id, s.assessment_id), messages, s.code
        )
        for stage, n in report["by_stage"].items():
            totals[stage] = totals.get(stage, 0) + n
        rows.append(
            {"user_id": student.id, "full_name": student.full_name, "task_id": s.task_id, **report}
        )
    return {"by_stage": totals, "students": rows}


# Research export to Excel (ФТ-49) ------------------------------------------------------------


FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _cell(value: Any) -> Any:
    """Text that a spreadsheet would run as a formula is kept as text (CSV/formula injection;
    e.g. a model name returned by the LLM provider)."""
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


@router.get("/assessments/{assessment_id}/export.xlsx")
def export_xlsx(
    assessment_id: int, request: Request, user: CurrentUser, db: DbSession, settings: AppSettings
) -> Response:
    """The same pseudonymous rows as the CSV, plus the agreement statistics."""
    assessment = _assessment_for_teacher(db, user, assessment_id)
    rows = _rows(db, assessment, settings.jwt_secret)
    book = Workbook()
    book.remove(book.worksheets[0])  # the empty default sheet
    sheet = book.create_sheet("grades")
    columns = list(rows[0]) if rows else ["student", "task_id"]
    sheet.append(columns)
    for row in rows:
        sheet.append([_cell(row[c]) for c in columns])
    stats = book.create_sheet("agreement")
    stats.append(["measure", "n", "pearson", "spearman", "mae", "md"])
    graded = [r for r in rows if r["ai_S"] != ""]
    for key in ("S", "F", "E", "A"):
        result = agreement(
            [float(r[f"ai_{key}"]) for r in graded], [float(r[f"teacher_{key}"]) for r in graded]
        )
        stats.append(
            [key, result["n"], result["pearson"], result["spearman"], result["mae"], result["md"]]
        )
    stats.append([])
    stats.append(["baseline (spec 16, Claude)", "", 0.647, "", 7.34, ""])
    buffer = io.BytesIO()
    book.save(buffer)
    audit.record(
        db,
        AuditAction.DATA_EXPORTED,
        actor_id=user.id,
        object_type="assessment",
        object_id=assessment_id,
        ip=client_ip(request),
        details={"rows": len(rows), "format": "xlsx"},
    )
    db.commit()
    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="assessment-{assessment_id}.xlsx"'},
    )
