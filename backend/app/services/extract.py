"""Text extraction from teaching materials (ФТ-06), keeping chapter and page (ФТ-35).

Runs in the separate `extractor` worker, which has no internet and no access to the
application's web tier (spec 14.1: text is extracted in an isolated processor).
Scanned PDF pages, which have no text layer, are recognised with Tesseract in Kazakh and
Russian (ФТ-37).
"""

import io
import posixpath
import re
import resource
import shutil
import subprocess  # nosec B404
import tempfile
import time
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from defusedxml import ElementTree

MAX_PAGES = 3000
MAX_UNPACKED = 100 * 1024 * 1024  # DOCX, PPTX, EPUB after unpacking
MAX_MEMBER = 20 * 1024 * 1024  # one EPUB chapter
CHUNK_CHARS = 1200
CHUNK_OVERLAP = 150
# A PDF page with less text than this is treated as a scan (ФТ-37).
MIN_PAGE_TEXT = 20
OCR_LANGUAGES = "kaz+rus"
OCR_MAX_PIXELS = 3000  # longest side of a rendered page: a huge page cannot exhaust memory
OCR_TIMEOUT_S = 60  # per tool and page
OCR_BUDGET_S = 420  # all OCR of one book: stays inside the 600 s material job
OCR_MAX_PAGES = 200
OCR_MEMORY = 1024**3

# pdf bytes, 1-based page numbers → recognised text per page
Ocr = Callable[[bytes, list[int]], dict[int, str]]


class ExtractError(ValueError):
    pass


@dataclass(frozen=True)
class Page:
    chapter: str
    page: int | None  # None when the format has no pages (DOCX, EPUB)
    text: str


@dataclass(frozen=True)
class Piece:
    chapter: str
    page: int | None
    text: str


# File checks ---------------------------------------------------------------------------

EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".epub": "application/epub+zip",
}


def detect_type(filename: str, data: bytes) -> str:
    """MIME type from the content, which must agree with the extension."""
    extension = posixpath.splitext(filename.lower())[1]
    if extension not in EXTENSIONS:
        raise ExtractError("unsupported_type")
    expected = EXTENSIONS[extension]
    if extension == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ExtractError("type_mismatch")
        return expected
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        names = set(archive.namelist())
    except zipfile.BadZipFile:
        raise ExtractError("type_mismatch") from None
    # Zip bombs: the declared sizes are checked before anything is unpacked.
    if sum(i.file_size for i in archive.infolist()) > MAX_UNPACKED:
        raise ExtractError("too_large_unpacked")
    try:
        if extension == ".epub":
            if archive.getinfo("mimetype").file_size > 64:
                raise ExtractError("type_mismatch")
            if archive.read("mimetype").strip() != b"application/epub+zip":
                raise ExtractError("type_mismatch")
            return expected
    except KeyError:
        raise ExtractError("type_mismatch") from None
    marker = "word/document.xml" if extension == ".docx" else "ppt/presentation.xml"
    if marker not in names or "[Content_Types].xml" not in names:
        raise ExtractError("type_mismatch")
    return expected


# Extraction ----------------------------------------------------------------------------


def _ocr_limits() -> None:
    """Memory and CPU limits of the poppler and tesseract processes."""
    resource.setrlimit(resource.RLIMIT_AS, (OCR_MEMORY, OCR_MEMORY))
    resource.setrlimit(resource.RLIMIT_CPU, (OCR_TIMEOUT_S, OCR_TIMEOUT_S))


def tesseract_ocr(data: bytes, numbers: list[int]) -> dict[int, str]:
    """Render each page with poppler and recognise it with Tesseract (kaz+rus). Missing
    tools, a failed page or the spent time budget give no text for the remaining pages
    instead of failing the book."""
    out: dict[int, str] = {}
    deadline = time.monotonic() + OCR_BUDGET_S
    pdftoppm, tesseract = shutil.which("pdftoppm"), shutil.which("tesseract")
    if pdftoppm is None or tesseract is None:
        return out
    with tempfile.TemporaryDirectory() as tmp:
        pdf = Path(tmp) / "book.pdf"
        pdf.write_bytes(data)
        for number in numbers:
            if time.monotonic() > deadline:
                break
            image = Path(tmp) / f"page{number}"
            try:
                subprocess.run(  # noqa: S603  # nosec B603
                    [
                        pdftoppm,
                        "-f",
                        str(number),
                        "-l",
                        str(number),
                        "-scale-to",
                        str(OCR_MAX_PIXELS),
                        "-gray",
                        "-png",
                        "-singlefile",
                        str(pdf),
                        str(image),
                    ],
                    check=True,
                    capture_output=True,
                    timeout=OCR_TIMEOUT_S,
                    preexec_fn=_ocr_limits,  # nosec B602
                )
                result = subprocess.run(  # noqa: S603  # nosec B603
                    [tesseract, f"{image}.png", "stdout", "-l", OCR_LANGUAGES],
                    check=True,
                    capture_output=True,
                    timeout=OCR_TIMEOUT_S,
                    preexec_fn=_ocr_limits,  # nosec B602
                )
            except (OSError, subprocess.SubprocessError):
                continue
            finally:
                Path(f"{image}.png").unlink(missing_ok=True)  # do not keep page images
            out[number] = result.stdout.decode("utf-8", errors="replace")
    return out


def extract_pdf(
    data: bytes, ocr: Ocr | None = None, max_ocr_pages: int = OCR_MAX_PAGES
) -> list[Page]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > MAX_PAGES:
            raise ExtractError("too_many_pages")
        chapters = _pdf_outline(reader)
        pages = []
        chapter = ""
        for number, page in enumerate(reader.pages, start=1):
            chapter = chapters.get(number, chapter)
            pages.append(Page(chapter, number, page.extract_text() or ""))
        scanned = [p.page for p in pages if p.page and len(p.text.strip()) < MIN_PAGE_TEXT]
        if ocr is not None and scanned:
            texts = ocr(data, scanned[:max_ocr_pages])
            pages = [
                Page(p.chapter, p.page, texts[p.page]) if p.page in texts else p for p in pages
            ]
        return pages
    except PdfReadError as exc:
        raise ExtractError(f"bad_pdf: {exc}") from None


def _pdf_outline(reader: object) -> dict[int, str]:
    """First page of each top-level outline entry → its title."""
    result: dict[int, str] = {}
    try:
        for item in reader.outline:  # type: ignore[attr-defined]
            if isinstance(item, list):
                continue
            page = reader.get_destination_page_number(item) + 1  # type: ignore[attr-defined]
            result.setdefault(page, str(item.title)[:300])
    except Exception:  # a broken outline only loses chapter names
        return {}
    return result


def extract_docx(data: bytes) -> list[Page]:
    import docx

    document = docx.Document(io.BytesIO(data))
    pages: list[Page] = []
    chapter = ""
    buffer: list[str] = []
    for paragraph in document.paragraphs:
        style = (paragraph.style.name if paragraph.style is not None else "").lower()
        if style.startswith(("heading 1", "heading 2", "title")) and paragraph.text.strip():
            if buffer:
                pages.append(Page(chapter, None, "\n".join(buffer)))
                buffer = []
            chapter = paragraph.text.strip()[:300]
        elif paragraph.text.strip():
            buffer.append(paragraph.text)
    if buffer:
        pages.append(Page(chapter, None, "\n".join(buffer)))
    return pages


def extract_pptx(data: bytes) -> list[Page]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    pages = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
        ]
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        pages.append(Page(title[:300], number, "\n".join(texts)))
    return pages


class _TextOf(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.title = ""
        self._in_heading = False
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style"):
            self._skip += 1
        if tag in ("h1", "h2") and not self.title:
            self._in_heading = True

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style") and self._skip:
            self._skip -= 1
        if tag in ("h1", "h2"):
            self._in_heading = False
        if tag in ("p", "div", "br", "li", "h1", "h2", "h3", "pre"):
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        if self._in_heading:
            self.title += data
        self.parts.append(data)


def extract_epub(data: bytes) -> list[Page]:
    archive = zipfile.ZipFile(io.BytesIO(data))
    container = ElementTree.fromstring(archive.read("META-INF/container.xml"))
    rootfile = next(e for e in container.iter() if e.tag.endswith("rootfile"))
    opf_path = str(rootfile.get("full-path"))
    opf = ElementTree.fromstring(archive.read(opf_path))
    manifest = {
        e.get("id"): e.get("href") for e in opf.iter() if e.tag.endswith("item") and e.get("href")
    }
    base = posixpath.dirname(opf_path)
    pages = []
    for itemref in (e for e in opf.iter() if e.tag.endswith("itemref")):
        href = manifest.get(itemref.get("idref"))
        if not href:
            continue
        path = posixpath.normpath(posixpath.join(base, href))
        if path.startswith(".."):
            continue
        parser = _TextOf()
        if archive.getinfo(path).file_size > MAX_MEMBER:
            raise ExtractError("chapter_too_large")
        parser.feed(archive.read(path).decode("utf-8", errors="replace"))
        text = re.sub(r"\n\s*\n+", "\n\n", "".join(parser.parts)).strip()
        if text:
            pages.append(Page(parser.title.strip()[:300], None, text))
    return pages


EXTRACTORS = {
    "application/pdf": extract_pdf,
    EXTENSIONS[".docx"]: extract_docx,
    EXTENSIONS[".pptx"]: extract_pptx,
    EXTENSIONS[".epub"]: extract_epub,
}


def extract(mime: str, data: bytes, ocr: Ocr | None = None) -> list[Page]:
    try:
        if mime == "application/pdf":
            return extract_pdf(data, ocr)
        return EXTRACTORS[mime](data)
    except ExtractError:
        raise
    except Exception as exc:  # untrusted file: any parser failure is a bad file
        raise ExtractError(f"cannot_extract: {type(exc).__name__}") from None


def split(pages: list[Page]) -> list[Piece]:
    """Chunks of about CHUNK_CHARS that never cross a page or chapter boundary."""
    pieces = []
    for page in pages:
        text = re.sub(r"[ \t]+", " ", page.text).strip()
        start = 0
        while start < len(text):
            end = min(len(text), start + CHUNK_CHARS)
            if end < len(text):
                cut = text.rfind(" ", start + CHUNK_CHARS // 2, end)
                end = cut if cut > start else end
            pieces.append(Piece(page.chapter, page.page, text[start:end].strip()))
            if end >= len(text):
                break
            start = max(end - CHUNK_OVERLAP, start + 1)
    return [p for p in pieces if p.text]
