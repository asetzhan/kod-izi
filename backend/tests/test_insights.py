"""Stage 6b, teacher analytics: ФТ-04, ФТ-05, ФТ-17, ФТ-20, ФТ-26, ФТ-28, ФТ-33, ФТ-38, ФТ-49."""

import io
from datetime import UTC, datetime, timedelta
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.enums import ActivityStage
from app.models import AiGrade, AiMessage, AppSetting, Submission, TraceEvent
from app.services.analytics import (
    SOLVE_RATE_KEY,
    error_kind,
    genai_report,
    label_stages,
    process_report,
)
from app.services.pipeline import run_pipeline
from tests.conftest import School, login
from tests.fake_judge import FakeJudge
from tests.llm_helpers import RecordingProvider
from tests.workspace_helpers import SOLUTION, _open, _submit, fast_analysis

LEVELS = {
    c: "high"
    for c in (
        "understanding",
        "algorithm",
        "correctness",
        "code_quality",
        "debugging",
        "explanation",
    )
}
T0 = datetime(2026, 11, 1, 10, 0, tzinfo=UTC)


def _event(kind: str, minute: float, **payload: Any) -> TraceEvent:
    return TraceEvent(
        kind=kind, payload=payload, server_ts=T0 + timedelta(minutes=minute), flagged=False
    )


def _confirm(client: TestClient, ws: dict[str, Any], submission_id: int) -> Any:
    response = client.post(
        f"/api/submissions/{submission_id}/teacher-grade",
        headers=ws["teacher"],
        json={
            "levels": LEVELS,
            "components": {"F": 100, "E": 60, "A": 60},
            "feedback": "ok",
            "reason": "Оқытушы деңгейлерді нақтылады.",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


# Activity stages and process metrics (ФТ-26, ФТ-33), pure ----------------------------------


SESSION = [
    _event("session_start", 0),
    _event("focus_lost", 1),
    _event("snapshot", 3, code="// read n\n// print parity"),
    _event("snapshot", 6, code="int main() { int n; cin >> n }"),
    _event("compile", 7, compile_output="main.cpp:1: error: expected ';' before '}'"),
    _event("snapshot", 8, code="int main() { int n; cin >> n; }"),
    _event("ai_request", 8.5, question="?"),
    _event("compile", 9, compile_output="main.cpp:1: error: expected ';' before 'return'"),
    _event("snapshot", 10, code="fixed"),
    _event("test", 12, compile_output="", per_test=["OK"]),
    _event("submit", 13),
    _event("verdict", 13.5, per_test=["OK", "OK"]),
    _event("explanation", 14, text="n % 2"),
    _event("session_closed", 15),
]


def test_ft26_trace_events_are_labelled_by_activity_stage() -> None:
    stages = label_stages(SESSION)
    assert stages == [
        None,
        ActivityStage.UNDERSTAND,
        ActivityStage.PLAN,  # comments before the first run: a plan in the editor
        ActivityStage.CODE,
        ActivityStage.TEST,
        ActivityStage.DEBUG,  # typing after a compile error
        ActivityStage.DEBUG,  # the AI request takes the student's stage
        ActivityStage.TEST,
        ActivityStage.DEBUG,
        ActivityStage.TEST,
        ActivityStage.TEST,
        ActivityStage.TEST,
        ActivityStage.EXPLAIN,
        None,
    ]


def test_ft33_process_metrics_errors_repetition_and_fix_time() -> None:
    report = process_report(SESSION)
    assert report["compile_errors"] == 2
    assert report["compile_error_kinds"] == {"missing_semicolon": 2}
    assert report["repeated_compile_errors"] == 1
    assert report["fix_cycles"] == 1 and report["fix_minutes"] == 5.0  # 7th → 12th minute
    assert report["attempts"] == 4  # two compile errors, one test run, one submission
    assert report["stage_events"]["debug"] == 3 and report["minutes"] == 15.0
    assert error_kind("error: 'x' was not declared in this scope") == "undeclared"
    assert error_kind("something odd") == "other"


def test_ft38_genai_use_by_stage_and_similarity() -> None:
    messages = [AiMessage(answer="int main() { int n; cin >> n; }", blocked=False, question="?")]
    messages[0].created_at = T0
    report = genai_report(SESSION, messages, "int main() { int n; cin >> n; }")
    assert report["by_stage"] == {"debug": 1}
    assert report["requests"] == 1 and report["max_similarity_to_final_code"] == 1.0


# Through the API -----------------------------------------------------------------------------


def test_ft26_ft28_trace_for_replay_carries_stages(client: TestClient, ws: dict[str, Any]) -> None:
    sid = _open(client, ws)
    client.post(
        f"/api/sessions/{sid}/events",
        headers=ws["student"],
        json={
            "events": [
                {
                    "seq": 1,
                    "kind": "snapshot",
                    "client_ts": datetime.now(UTC).isoformat(),
                    "payload": {"code": SOLUTION},
                }
            ]
        },
    )
    client.post(f"/api/sessions/{sid}/test", headers=ws["student"], json={"code": SOLUTION})
    trace = client.get(f"/api/sessions/{sid}/trace", headers=ws["teacher"]).json()
    by_kind = {e["kind"]: e["stage"] for e in trace["events"]}
    assert by_kind == {"session_start": None, "snapshot": "code", "test": "test"}
    snapshot = next(e for e in trace["events"] if e["kind"] == "snapshot")
    assert snapshot["payload"]["code"] == SOLUTION  # the replay rebuilds the code from these


def test_ft04_monitoring_board_for_the_teacher(
    client: TestClient, ws: dict[str, Any], school: School
) -> None:
    good = _submit(client, ws, _open(client, ws))
    _confirm(client, ws, good["id"])
    _submit(client, ws, _open(client, ws), "// fake:WA\n" + SOLUTION)
    board = client.get(f"/api/courses/{ws['course']['id']}/monitoring", headers=ws["teacher"])
    assert board.status_code == 200
    data = board.json()
    assert data["criteria"]["correctness"] == {"high": 1}
    assert data["difficulties"]["verdicts"] == {"WA": 1}
    [row] = data["students"]
    assert row["full_name"] == school["student_a"].full_name
    assert row["submissions"] == 2 and row["solved"] == 1
    other = login(client, school["teacher_b"].email)
    assert (
        client.get(f"/api/courses/{ws['course']['id']}/monitoring", headers=other).status_code
        == 403
    )


def test_ft05_individual_task_is_worked_on_and_graded(
    client: TestClient, ws: dict[str, Any], school: School, db: Session
) -> None:
    url = f"/api/courses/{ws['course']['id']}/assignments"
    body = {"task_id": ws["task"]["id"], "user_id": school["student_a"].id, "note": "Қайталау"}
    created = client.post(url, headers=ws["teacher"], json=body)
    assert created.status_code == 201 and created.json()["genai_mode"] == "hints_only"
    assert client.post(url, headers=ws["teacher"], json=body).json()["detail"] == "already_assigned"
    # Not to a student the teacher does not teach.
    stranger = {**body, "user_id": school["student_b"].id}
    assert client.post(url, headers=ws["teacher"], json=stranger).status_code == 403

    opened = client.post(
        "/api/sessions",
        headers=ws["student"],
        json={"task_id": ws["task"]["id"], "assessment_id": None},
    )
    assert opened.status_code == 201 and opened.json()["genai_mode"] == "hints_only"
    sid = opened.json()["session_id"]
    no_text = client.post(
        f"/api/sessions/{sid}/submit", headers=ws["student"], json={"code": SOLUTION}
    )
    assert no_text.json()["detail"] == "explanation_required"
    submitted = _submit(client, ws, sid)
    assert db.scalar(select(AiGrade).where(AiGrade.submission_id == submitted["id"])) is not None
    mine = client.get("/api/my/assignments", headers=ws["student"]).json()
    assert mine[0]["note"] == "Қайталау" and mine[0]["last_verdict"] == "OK"
    other = login(client, school["student_b"].email)
    assert (
        client.post(
            "/api/sessions",
            headers=other,
            json={"task_id": ws["task"]["id"], "assessment_id": None},
        ).status_code
        == 403
    )


def test_ft17_graded_works_become_examples_for_the_llm(
    client: TestClient, ws: dict[str, Any], db: Session, judge: FakeJudge
) -> None:
    first = _submit(client, ws, _open(client, ws))
    url = f"/api/submissions/{first['id']}/examples"
    early = client.post(url, headers=ws["teacher"], json={"criterion_code": "code_quality"})
    assert early.status_code == 409 and early.json()["detail"] == "not_confirmed"
    _confirm(client, ws, first["id"])
    example = client.post(
        url, headers=ws["teacher"], json={"criterion_code": "code_quality", "note": "Таза код"}
    ).json()
    assert example["level"] == "high"  # the teacher's level, not the AI's
    listed = client.get(f"/api/courses/{ws['course']['id']}/rubric/examples", headers=ws["teacher"])
    assert [e["criterion_code"] for e in listed.json()] == ["code_quality"]
    assert client.post(url, headers=ws["student"], json={"criterion_code": "x"}).status_code == 403

    second = _submit(client, ws, _open(client, ws))
    provider = RecordingProvider()
    submission = db.get(Submission, second["id"])
    assert submission is not None
    submission.status = "queued"
    run_pipeline(db, judge, provider, submission, analyser=fast_analysis)
    [request] = [r for r in provider.requests if r.prompt_version.startswith("assessment-")]
    assert request.prompt_version.endswith("+examples")
    assert "Таза код" in request.data_blocks["graded_examples"]
    assert "graded_examples" in request.system
    graded = db.scalars(select(AiGrade).where(AiGrade.submission_id == second["id"])).all()
    assert graded[-1].prompt_version.endswith("+examples")


def test_ft20_tasks_whose_solve_rate_does_not_fit_the_level_are_flagged(
    client: TestClient, ws: dict[str, Any], db: Session
) -> None:
    db.add(AppSetting(key=SOLVE_RATE_KEY, value={"tolerance": 0.25, "min_students": 1}))
    db.flush()
    _submit(client, ws, _open(client, ws), "// fake:WA\n" + SOLUTION)
    data = client.get(
        f"/api/courses/{ws['course']['id']}/solve-rates", headers=ws["teacher"]
    ).json()
    [row] = data["tasks"]
    assert row["level"] == "I" and row["rate"] == 0.0 and row["expected"] == 0.9894
    assert row["flagged"] and row["direction"] == "harder"
    _submit(client, ws, _open(client, ws))
    [row] = client.get(
        f"/api/courses/{ws['course']['id']}/solve-rates", headers=ws["teacher"]
    ).json()["tasks"]
    assert row["rate"] == 1.0 and not row["flagged"]


def test_ft33_ft38_assessment_reports(client: TestClient, ws: dict[str, Any]) -> None:
    _submit(client, ws, _open(client, ws))
    process = client.get(
        f"/api/assessments/{ws['assessment']['id']}/process", headers=ws["teacher"]
    ).json()
    assert process[0]["attempts"] >= 1 and "compile_error_kinds" in process[0]
    genai = client.get(f"/api/assessments/{ws['assessment']['id']}/genai", headers=ws["teacher"])
    assert genai.status_code == 200 and genai.json()["students"][0]["requests"] == 0
    assert (
        client.get(
            f"/api/assessments/{ws['assessment']['id']}/genai", headers=ws["student"]
        ).status_code
        == 403
    )


def test_ft49_research_export_to_excel(
    client: TestClient, ws: dict[str, Any], school: School
) -> None:
    result = _submit(client, ws, _open(client, ws))
    _confirm(client, ws, result["id"])
    response = client.get(
        f"/api/assessments/{ws['assessment']['id']}/export.xlsx", headers=ws["teacher"]
    )
    assert response.status_code == 200 and "attachment" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["grades", "agreement"]
    rows = list(book["grades"].values)
    assert rows[0][0] == "student" and str(rows[1][0]).startswith("S-")
    text = " ".join(str(v) for row in rows for v in row)
    assert school["student_a"].full_name not in text and school["student_a"].email not in text
    stats = {r[0]: r for r in book["agreement"].values if r and r[0]}
    assert stats["S"][1] == 1
    assert (
        client.get(
            f"/api/assessments/{ws['assessment']['id']}/export.xlsx", headers=ws["student"]
        ).status_code
        == 403
    )
